#!/usr/bin/env python3
"""
Standalone script to import EM-DAT data from Excel file
"""
import sys
import os

# Add the disaster_management directory to the path
sys.path.insert(0, '/home/mahmoud/disaster_management')

from evaluation_service.modules.historical.importer import HistoricalImporter

if __name__ == '__main__':
    excel_path = '/home/mahmoud/public_emdat_incl_hist_2026-05-25.xlsx'
    limit = 1000  # Import first 1000 records with coordinates
    
    print(f"Importing EM-DAT data from {excel_path}...")
    print(f"Limit: {limit} records with coordinates")
    
    # First, check how many rows have coordinates
    try:
        import openpyxl
        wb = openpyxl.load_workbook(excel_path, read_only=True)
        sheet = wb.active
        headers = list(sheet.iter_rows(max_row=1, values_only=True))[0]
        header_map = {h: i for i, h in enumerate(headers)}
        
        idx_lat = header_map.get('Latitude')
        idx_lng = header_map.get('Longitude')
        
        print(f"\nColumn indices: Latitude={idx_lat}, Longitude={idx_lng}")
        
        # Count rows with coordinates
        rows_with_coords = 0
        total_rows = 0
        for row in sheet.iter_rows(min_row=2, values_only=True):
            total_rows += 1
            lat = row[idx_lat] if idx_lat is not None else None
            lng = row[idx_lng] if idx_lng is not None else None
            if lat and lng:
                rows_with_coords += 1
            if total_rows <= 5:
                print(f"Row {total_rows}: lat={lat}, lng={lng}")
        
        print(f"\nTotal rows: {total_rows}")
        print(f"Rows with coordinates: {rows_with_coords}")
        wb.close()
    except Exception as e:
        print(f"Error checking Excel file: {e}")
    
    importer = HistoricalImporter()
    success_count, error_count = importer.import_from_excel(excel_path, limit)
    
    print(f"\nImport complete:")
    print(f"  Success: {success_count} disasters")
    print(f"  Errors: {error_count}")
